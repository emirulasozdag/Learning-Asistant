from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape

from app.ingestion.types import RawSegment


def parse_pptx(path: Path) -> list[RawSegment]:
    """PPTX → slayt başına bir segment; slayt numarası `page` olarak tutulur.

    Başlık, metin kutuları, tablolar, gruplar ve konuşmacı notları okunur.
    """
    presentation = Presentation(str(path))
    segments: list[RawSegment] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        margins = {}
        warnings = []
        for shape in slide.shapes:
            extracted = _shape_lines(shape)
            if shape.is_placeholder and shape.placeholder_format.type in {
                PP_PLACEHOLDER.HEADER,
                PP_PLACEHOLDER.FOOTER,
                PP_PLACEHOLDER.DATE,
                PP_PLACEHOLDER.SLIDE_NUMBER,
            }:
                if shape.top + shape.height <= presentation.slide_height * 0.1 or (
                    shape.top >= presentation.slide_height * 0.9
                ):
                    start = len("\n".join(lines).split("\n")) if lines else 0
                    for offset in range(len("\n".join(extracted).split("\n"))):
                        margins[start + offset] = (
                            shape.left / presentation.slide_width,
                            shape.top / presentation.slide_height,
                        )
            lines.extend(extracted)
            if _has_visual_content(shape):
                warnings.append(
                    f"Slayt {number}: Tablo, formül, şekil veya görselin yerleşimi metinde tam "
                    "temsil edilemeyebilir; kaynak slaytı kontrol edin."
                )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(notes)
        segments.append(
            RawSegment(
                text="\n".join(lines),
                page=number,
                margin_lines=margins,
                warnings=list(dict.fromkeys(warnings)),
            )
        )
    return segments


def _shape_lines(shape: BaseShape) -> list[str]:
    if isinstance(shape, GroupShape):
        return [line for child in shape.shapes for line in _shape_lines(child)]
    if shape.has_text_frame:
        return [
            "    " * paragraph.level + paragraph.text for paragraph in shape.text_frame.paragraphs
        ]
    if getattr(shape, "has_table", False):
        return [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
    return []


def _has_visual_content(shape: BaseShape) -> bool:
    if isinstance(shape, GroupShape):
        return any(_has_visual_content(child) for child in shape.shapes)
    if getattr(shape, "has_table", False) and any(
        cell.is_merge_origin or cell.is_spanned for row in shape.table.rows for cell in row.cells
    ):
        return True
    return shape.shape_type in {
        MSO_SHAPE_TYPE.PICTURE,
        MSO_SHAPE_TYPE.CHART,
        MSO_SHAPE_TYPE.DIAGRAM,
        MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
        MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
        MSO_SHAPE_TYPE.AUTO_SHAPE,
        MSO_SHAPE_TYPE.LINE,
        MSO_SHAPE_TYPE.FREEFORM,
    } or any(
        element.tag.endswith("}oMath") or element.get("baseline", "0") != "0"
        for element in shape.element.iter()
    )
