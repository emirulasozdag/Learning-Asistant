"""Ortak test yardımcıları.

Veritabanı gerektiren testler yalnızca TEST_DATABASE_URL tanımlıysa çalışır (CI'da tanımlı).
Lokal: docker compose ile açılan PostgreSQL'de `learning_assistant_test` veritabanını oluşturup
  TEST_DATABASE_URL=postgresql+psycopg://learning:learning@localhost:5432/learning_assistant_test
ile pytest çalıştırın. Bu veritabanının tabloları her testte boşaltılır.
"""

import os
from collections.abc import Callable, Iterator
from datetime import UTC, datetime
from functools import cache
from pathlib import Path

import pymupdf
import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import create_engine, make_url, text
from sqlalchemy.orm import Session, sessionmaker

from app.api.deps import get_session_factory
from app.auth.passwords import hash_password
from app.auth.sessions import create_session
from app.core.config import settings
from app.db.models import Course, Enrollment, User
from app.db.session import get_db
from app.main import app
from app.schemas.common import Role

BACKEND_DIR = Path(__file__).resolve().parents[1]
TEST_DATABASE_URL = os.environ.get("TEST_DATABASE_URL")

# Testlerde kullanılan örnek ders içeriği: (başlık, gövde) — sayfa / slayt başına bir tane.
SAMPLE_PAGES = [
    (
        "Ikili Arama",
        "Ikili arama sirali bir dizide hedef degeri arar. "
        "Her adimda ortadaki eleman hedefle karsilastirilir.",
    ),
    (
        "Neden yariyi atabiliriz",
        "Dizi sirali oldugu icin hedef ortadaki elemandan kucukse sag yari elenir "
        "ve arama sol yarida surer.",
    ),
    ("Karmasiklik", "Zaman karmasikligi O(log n) olur. Dogrusal arama ise O(n) zaman alir."),
]


def make_pdf(path: Path, pages: list[tuple[str, str]], header: str | None = None) -> Path:
    document = pymupdf.open()
    for number, (title, body) in enumerate(pages, start=1):
        page = document.new_page()
        if header:
            page.insert_text((50, 40), header, fontsize=8)
        page.insert_textbox(pymupdf.Rect(50, 70, 545, 110), title, fontsize=16)
        page.insert_textbox(pymupdf.Rect(50, 120, 545, 760), body, fontsize=11)
        page.insert_text((290, 820), str(number), fontsize=8)
    document.save(path)
    return path


def make_pptx(path: Path, pages: list[tuple[str, str]], notes: bool = True) -> Path:
    presentation = Presentation()
    for title, body in pages:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = f"Not: {title}"
    presentation.save(str(path))
    return path


@pytest.fixture(autouse=True)
def _fast_embeddings(monkeypatch: pytest.MonkeyPatch) -> None:
    """Testler gerçek modeli indirmez: kelime tabanlı sahte embedder kullanılır."""
    monkeypatch.setattr(settings, "embedding_provider", "fake")


@pytest.fixture(autouse=True)
def _claim_check_off(monkeypatch: pytest.MonkeyPatch) -> None:
    """İddia kontrolü (#10) yalnızca onu sınayan testlerde açılır (`test_claims.py`); diğer
    testler sahte üreticiyle alıntı doğrulamasını sınar. Yerel .env'deki judge ayarı da sızmaz."""
    monkeypatch.setattr(settings, "claim_check", "off")
    monkeypatch.setattr(settings, "judge_provider", "")
    monkeypatch.setattr(settings, "judge_model", "")


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    return make_pdf(tmp_path / "notlar.pdf", SAMPLE_PAGES, header="BLM2512 Veri Yapilari")


@pytest.fixture
def sample_pptx(tmp_path: Path) -> Path:
    return make_pptx(tmp_path / "slaytlar.pptx", SAMPLE_PAGES)


@pytest.fixture(scope="session")
def db_engine():
    if not TEST_DATABASE_URL:
        if os.environ.get("REQUIRE_DB_TESTS"):
            pytest.fail(
                "REQUIRE_DB_TESTS ayarlı ama TEST_DATABASE_URL yok: veritabanı testleri atlanamaz"
            )
        pytest.skip("TEST_DATABASE_URL tanımlı değil; veritabanı testleri atlandı")
    if not make_url(TEST_DATABASE_URL).database.endswith("_test"):
        # Tablolar her testte boşaltılır: geliştirme veritabanı yanlışlıkla silinmesin.
        pytest.fail("TEST_DATABASE_URL'deki veritabanı adı '_test' ile bitmeli")
    config = Config(str(BACKEND_DIR / "alembic.ini"))
    config.set_main_option("script_location", str(BACKEND_DIR / "migrations"))
    config.set_main_option("sqlalchemy.url", TEST_DATABASE_URL)
    command.upgrade(config, "head")
    engine = create_engine(TEST_DATABASE_URL)
    yield engine
    engine.dispose()


@pytest.fixture
def db(db_engine) -> Iterator[Session]:
    with db_engine.begin() as connection:
        connection.execute(
            text(
                "TRUNCATE messages, conversations, chunks, documents, enrollments, sessions, "
                "users, courses RESTART IDENTITY CASCADE"
            )
        )
    with session_factory(db_engine)() as session:
        yield session


def session_factory(engine) -> sessionmaker[Session]:
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


TEST_PASSWORD = "test-sifre-123"


@cache
def _test_password_hash() -> str:
    # argon2 bilerek yavaştır: testlerde bir kez hesaplanır.
    return hash_password(TEST_PASSWORD)


def make_user(
    db: Session,
    email: str,
    *,
    can_create_courses: bool = False,
    consented: bool = True,
    research_logging: bool = False,
) -> User:
    """Test kullanıcısı (şifresi TEST_PASSWORD). `consented`: güncel aydınlatma metni onaylı."""
    user = User(
        name=email.split("@")[0],
        email=email,
        password_hash=_test_password_hash(),
        can_create_courses=can_create_courses,
        research_logging=research_logging,
        policy_version=settings.privacy_policy_version if consented else None,
        consented_at=datetime.now(UTC) if consented else None,
    )
    db.add(user)
    db.commit()
    return user


def enroll(db: Session, user: User, course: Course, role: Role) -> Enrollment:
    enrollment = Enrollment(user_id=user.id, course_id=course.id, role=role)
    db.add(enrollment)
    db.commit()
    return enrollment


@pytest.fixture
def api(db: Session, db_engine, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """API'yi test veritabanına bağlar; LLM'siz ve geçici upload klasörlü."""
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    monkeypatch.setattr(settings, "cache_dir", tmp_path / "cache")
    monkeypatch.setattr(settings, "llm_api_key", "")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_session_factory] = lambda: session_factory(db_engine)
    yield
    app.dependency_overrides.clear()


@pytest.fixture
def client_for(api, db: Session) -> Callable[[User | None], TestClient]:
    """Verilen kullanıcının oturumuyla (None: oturumsuz) ayrı bir API istemcisi."""

    def make(user: User | None) -> TestClient:
        client = TestClient(app)
        if user is not None:
            client.cookies.set(settings.session_cookie_name, create_session(db, user))
        return client

    return make


@pytest.fixture
def user(db: Session) -> User:
    """Ders açabilen, aydınlatma metnini onaylamış kullanıcı; `course`'un eğitmeni."""
    return make_user(db, "egitmen@example.com", can_create_courses=True)


@pytest.fixture
def client(client_for, user: User) -> TestClient:
    """`user` olarak oturum açmış API istemcisi."""
    return client_for(user)


@pytest.fixture
def course(db: Session, user: User) -> Course:
    course = Course(code="BLM2512", name="Veri Yapilari ve Algoritmalar", language="tr")
    db.add(course)
    db.commit()
    enroll(db, user, course, "instructor")
    return course


def vtt_text(cues: list[tuple[float, float, str]]) -> str:
    """(başlangıç, bitiş, metin) listesinden geçerli bir WebVTT dosyası metni."""

    def stamp(seconds: float) -> str:
        total_ms = round(seconds * 1000)
        hours, rest = divmod(total_ms, 3_600_000)
        minutes, rest = divmod(rest, 60_000)
        return f"{hours:02}:{minutes:02}:{rest // 1000:02}.{rest % 1000:03}"

    blocks = [f"{stamp(start)} --> {stamp(end)}\n{text}" for start, end, text in cues]
    return "WEBVTT\n\n" + "\n\n".join(blocks) + "\n"


def make_vtt(path: Path, cues: list[tuple[float, float, str]]) -> Path:
    path.write_text(vtt_text(cues), encoding="utf-8")
    return path
