"""Materyal yükleme modülü: parser'lar, temizleme, parçalama (veritabanı gerektirmez)."""

from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.embeddings.fake import FakeEmbedder
from app.ingestion.chunking import chunk_segments
from app.ingestion.cleaning import clean_segments, clean_segments_with_stats
from app.ingestion.parsers.pdf_parser import parse_pdf
from app.ingestion.parsers.pptx_parser import parse_pptx
from app.ingestion.types import RawSegment
from tests.conftest import SAMPLE_PAGES, make_pdf, make_pptx

EMBEDDER = FakeEmbedder()


def test_parse_pdf_returns_one_segment_per_page_with_page_numbers(sample_pdf: Path) -> None:
    segments = parse_pdf(sample_pdf)

    assert [s.page for s in segments] == [1, 2, 3]
    assert "Ikili arama sirali" in segments[0].text
    assert "sag yari elenir" in segments[1].text


def test_parse_pdf_keeps_empty_pages_so_page_count_is_correct(tmp_path: Path) -> None:
    path = make_pdf(tmp_path / "bos.pdf", [("", ""), ("Baslik", "Icerik")])

    segments = parse_pdf(path)

    assert [s.page for s in segments] == [1, 2]
    assert segments[0].text.strip() == "1"


def test_parse_pptx_reads_title_body_and_notes_with_slide_numbers(sample_pptx: Path) -> None:
    segments = parse_pptx(sample_pptx)

    assert [s.page for s in segments] == [1, 2, 3]
    assert segments[1].text.startswith("Neden yariyi atabiliriz")
    assert "sag yari elenir" in segments[1].text
    assert "Not: Neden yariyi atabiliriz" in segments[1].text


def test_parse_pptx_handles_turkish_characters(tmp_path: Path) -> None:
    path = make_pptx(tmp_path / "tr.pptx", [("İkili Arama", "Sıralı dizide ağaç, şu, ğ")])

    assert "İkili Arama\nSıralı dizide ağaç, şu, ğ" in parse_pptx(path)[0].text


def test_cleaning_removes_repeated_header_and_page_numbers(sample_pdf: Path) -> None:
    segments = parse_pdf(sample_pdf) + parse_pdf(sample_pdf)  # 6 sayfa, her birinde aynı üst bilgi

    cleaned = clean_segments(segments)

    assert len(cleaned) == 6
    for segment in cleaned:
        assert "BLM2512" not in segment.text
        assert not any(line.strip().isdigit() for line in segment.text.split("\n"))


def test_cleaning_keeps_section_title_repeated_on_some_pages() -> None:
    segments = [RawSegment(text=f"Ikili Arama\nicerik {i}", page=i) for i in range(1, 4)]
    segments += [RawSegment(text=f"Siralama\nicerik {i}", page=i) for i in range(4, 7)]

    cleaned = clean_segments(segments)

    assert all(s.text.split("\n")[0] in {"Ikili Arama", "Siralama"} for s in cleaned)


def test_cleaning_drops_empty_pages_without_rewriting_source_punctuation() -> None:
    segments = [
        RawSegment(text="   \n\n", page=1),
        RawSegment(text="hi = n ­ 1\nalgo-\nritma  ve   kelime", page=2),
    ]

    cleaned = clean_segments(segments)

    assert len(cleaned) == 1
    assert cleaned[0].page == 2
    assert cleaned[0].text == segments[1].text


def test_short_page_becomes_single_chunk_with_its_page() -> None:
    chunks = chunk_segments([RawSegment(text="Kisa bir sayfa.", page=7)], EMBEDDER)

    assert len(chunks) == 1
    assert chunks[0].page == 7
    assert chunks[0].text == "Kisa bir sayfa."


def test_long_page_is_split_under_token_limit_without_crossing_pages() -> None:
    sentence = "Ikili arama her adimda arama araligini yariya indirir. "
    long_page = RawSegment(text=(sentence * 120).strip(), page=3)
    next_page = RawSegment(text="Sonraki sayfa.", page=4)

    chunks = chunk_segments([long_page, next_page], EMBEDDER, max_tokens=100, overlap_tokens=20)

    page3 = [c for c in chunks if c.page == 3]
    assert len(page3) > 1
    assert all(EMBEDDER.count_passage_tokens(c.text) <= 100 for c in page3)
    assert chunks[-1].page == 4 and chunks[-1].text == "Sonraki sayfa."


def test_consecutive_chunks_overlap() -> None:
    sentences = [f"Cumle numarasi {i} burada biter." for i in range(60)]
    chunks = chunk_segments(
        [RawSegment(text=" ".join(sentences), page=1)], EMBEDDER, max_tokens=60, overlap_tokens=20
    )

    assert len(chunks) > 2
    for previous, current in zip(chunks, chunks[1:], strict=False):
        assert any(previous.text.split()[-n:] == current.text.split()[:n] for n in range(2, 21))


def test_sample_material_end_to_end_produces_one_chunk_per_page(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path / "a.pdf", SAMPLE_PAGES * 2, header="Ders Notlari")

    chunks = chunk_segments(clean_segments(parse_pdf(pdf)), EMBEDDER)

    assert [c.page for c in chunks] == [1, 2, 3, 4, 5, 6]
    assert all("Ders Notlari" not in c.text for c in chunks)


def test_cleaning_preserves_executable_python_and_numeric_output() -> None:
    code = "def answer():\n    if True:\n        return 42\n\n42\nx = n -\n1\n"
    cleaned = clean_segments([RawSegment(text=code, page=1)])[0]

    assert cleaned.text == code
    namespace = {}
    exec(compile(cleaned.text.split("\n\n")[0], "source", "exec"), namespace)
    assert namespace["answer"]() == 42


def test_pdf_furniture_is_removed_without_removing_repeated_body_or_numbers(tmp_path: Path) -> None:
    path = tmp_path / "layout.pdf"
    with pymupdf.open() as pdf:
        for number in range(1, 6):
            page = pdf.new_page()
            page.insert_text((50, 35), "Course header", fontsize=8)
            page.insert_text((50, 140), "def answer():\n    return 42", fontname="cour")
            page.insert_text((50, 220), "42\n1\n1\n-7")
            page.insert_text((50, 300), "value    frequency\n42       7", fontname="cour")
            page.insert_text((290, 820), str(number), fontsize=8)
        pdf.save(path)

    segments = parse_pdf(path)
    cleaned, stats = clean_segments_with_stats(segments)

    assert stats.repeated_lines_removed == 5
    # Page 1's footer is ambiguous with body occurrences of "1": conservatively retain it.
    assert stats.page_numbers_removed == 4
    assert all("Course header" not in segment.text for segment in cleaned)
    for segment in cleaned:
        assert "    return 42" in segment.text
        assert "42\n1\n1\n-7" in segment.text
        assert "value    frequency" in segment.text
        assert segment.warnings


def test_repeated_text_at_different_margin_positions_is_retained() -> None:
    segments = [
        RawSegment(text="Heading\ncontent", page=i, margin_lines={0: (0.1, i / 10)})
        for i in range(1, 5)
    ]
    assert [s.text for s in clean_segments(segments)] == [s.text for s in segments]


def test_pdf_partial_image_page_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "partial.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((50, 120), "Readable page")
        pdf.new_page()
        pdf.save(path)
    segments = parse_pdf(path)
    assert "OCR" in " ".join(segments[1].warnings)
    assert len(clean_segments(segments)) == 1


def test_repeated_monospace_code_in_top_margin_is_content(tmp_path: Path) -> None:
    path = tmp_path / "code-at-top.pdf"
    with pymupdf.open() as pdf:
        for number in range(1, 5):
            page = pdf.new_page()
            page.insert_text((50, 40), "return 42", fontname="cour")
            page.insert_text((50, 200), f"Example {number}")
        pdf.save(path)
    assert all("return 42" in segment.text for segment in clean_segments(parse_pdf(path)))


def test_pptx_keeps_nested_code_formula_and_table_cells(tmp_path: Path) -> None:
    path = tmp_path / "structured.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    text.text_frame.text = "def answer():\n    return 42\n42\nx = a - b"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
    for cell, value in zip(
        [cell for row in table.rows for cell in row.cells],
        ["value", "count", "42", "7"],
        strict=True,
    ):
        cell.text = value
    presentation.save(path)
    cleaned = clean_segments(parse_pptx(path))[0].text
    assert "def answer():\n    return 42\n42\nx = a - b" in cleaned
    assert "value | count\n42 | 7" in cleaned


def test_chunking_preserves_every_character_and_code_indentation() -> None:
    code = "def example():\n" + "    value += 42\n" * 120
    chunks = chunk_segments(
        [RawSegment(text=code, page=4)], EMBEDDER, max_tokens=40, overlap_tokens=0
    )
    assert "".join(chunk.text for chunk in chunks) == code
    assert all(chunk.page == 4 for chunk in chunks)
    assert all(EMBEDDER.count_passage_tokens(chunk.text) <= 40 for chunk in chunks)


def test_pptx_reports_merged_cells_and_spatial_math(tmp_path: Path) -> None:
    path = tmp_path / "unsupported-layout.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    shape.text_frame.text = "x"
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "2"
    run._r.get_or_add_rPr().set("baseline", "30000")
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged heading"
    table.cell(1, 0).text = "42"
    table.cell(1, 1).text = "7"
    presentation.save(path)
    segment = parse_pptx(path)[0]
    assert "42 | 7" in segment.text
    assert "x2" in segment.text
    assert segment.warnings


@pytest.mark.parametrize("budget,overlap", [(0, 0), (513, 0), (40, -1)])
def test_invalid_chunk_budgets_fail_clearly(budget: int, overlap: int) -> None:
    with pytest.raises(ValueError):
        chunk_segments([RawSegment(text="content", page=1)], EMBEDDER, budget, overlap)
